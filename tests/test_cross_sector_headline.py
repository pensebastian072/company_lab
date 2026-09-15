"""The cross-sector presentation, added 2026-08-25.

This change moves NOTHING in the scoring. What it does is put the answer to "is this
company good compared with companies in other sectors?" in the first screenful, because
E03 R1 measured 95.6% of the ranking's power as sector selection and the workbook led
with the column that cannot answer it.

So the tests here pin three things, all of them about honesty rather than layout:

  * the cross-sector columns exist and are populated on every path that builds a
    workbook, INCLUDING one rebuilt from a parquet written before these columns existed
    - a headline column that renders blank because an old file lacks it is the
    graceful-degradation failure this repo keeps hitting;
  * the note says which score the percentile was taken over, because it is
    `composite_strict` and the column sitting beside it is `score`, and they disagree;
  * the workbook never claims the change improved anything.
"""
from __future__ import annotations

import io

from clab.export import csv_export, xlsx_export
from clab.scoring import sector_neutral as sn


def _universe():
    """Two sectors, wide apart, with enough companies to clear both peer floors."""
    rows = []
    for i in range(8):
        rows.append({"ticker": f"IT{i}", "name": f"Tech {i}", "cik": f"1{i}",
                     "sector": "Information Technology", "sub_industry": "Software",
                     "composite_strict": 55 + i, "score": 55 + i,
                     "band": "INVESTABLE", "coverage": 0.93})
    for i in range(8):
        rows.append({"ticker": f"RE{i}", "name": f"Realty {i}", "cik": f"2{i}",
                     "sector": "Real Estate", "sub_industry": "REITs",
                     "composite_strict": 28 + i, "score": 28 + i,
                     "band": "INSUFFICIENT_DATA" if i < 6 else "WEAK",
                     "coverage": 0.71 if i < 6 else 0.84})
    return rows


def _book(rows, sheets=("rank", "sectors", "findings")):
    from openpyxl import load_workbook
    data = xlsx_export.to_bytes(rows, cards=[], sheets=sheets)
    return load_workbook(io.BytesIO(data))


def _text(ws):
    return "\n".join(str(c) for row in ws.iter_rows(values_only=True)
                     for c in row if c is not None)


# ------------------------------------------------------------------ the columns
def test_sector_rank_is_a_place_within_its_own_sector():
    rows = _universe()
    sn.neutralize(rows)
    best_re = next(r for r in rows if r["ticker"] == "RE7")
    worst_it = next(r for r in rows if r["ticker"] == "IT0")
    assert best_re["sector_rank"] == 1
    assert best_re["sector_rank_n"] == 8
    # the point of the whole change: last in IT outranks nothing, first in Real Estate
    # leads its sector, and the raw score says the opposite
    assert best_re["composite_strict"] < worst_it["composite_strict"]
    assert best_re["sector_rank"] < worst_it["sector_rank"]


def test_ties_share_a_rank_and_the_next_one_skips():
    rows = [{"ticker": t, "sector": "S", "sub_industry": "S", "composite_strict": v}
            for t, v in (("A", 90), ("B", 80), ("C", 80), ("D", 70), ("E", 60))]
    sn.neutralize(rows)
    assert [r["sector_rank"] for r in rows] == [1, 2, 2, 4, 5]


def test_sector_rank_has_no_peer_floor_because_it_is_a_count_not_a_statistic():
    """`sector_neutral_score` abstains under MIN_PEERS; "1 of 3" is still just true."""
    rows = [{"ticker": "A", "sector": "Tiny", "sub_industry": "Tiny",
             "composite_strict": 90},
            {"ticker": "B", "sector": "Tiny", "sub_industry": "Tiny",
             "composite_strict": 10}]
    sn.neutralize(rows)
    assert rows[0]["sector_rank"] == 1 and rows[0]["sector_rank_n"] == 2
    assert rows[0]["sector_neutral_score"] is None, "the percentile must still abstain"


def test_every_row_records_which_score_column_the_percentile_ranked():
    rows = _universe()
    sn.neutralize(rows)
    assert all(r["sector_neutral_score_col"] == "composite_strict" for r in rows)


def test_an_unscored_company_gets_no_rank():
    rows = _universe()
    rows.append({"ticker": "NONE", "sector": "Real Estate", "sub_industry": "REITs",
                 "composite_strict": None})
    sn.neutralize(rows)
    n = next(r for r in rows if r["ticker"] == "NONE")
    assert n["sector_rank"] is None and n["sector_rank_n"] == 0


# ------------------------------------------------------------------ ensure()
def test_ensure_is_a_no_op_once_the_columns_are_there():
    rows = _universe()
    sn.neutralize(rows)
    assert sn.ensure(rows) is False


def test_ensure_backfills_a_row_set_from_an_older_parquet():
    """The real case: a parquet written before sector_rank existed."""
    rows = _universe()
    sn.neutralize(rows)
    for r in rows:
        del r["sector_rank"]
    assert sn.ensure(rows) is True
    assert all(r["sector_rank"] is not None for r in rows)


# ------------------------------------------------------------------ the CSV
def test_the_cross_sector_answer_is_in_the_first_screenful():
    head = csv_export.COLUMNS[:12]
    assert "sector_rank" in head and "sector_neutral_score" in head
    assert csv_export.COLUMNS.index("score") < csv_export.COLUMNS.index("sector_rank")


def test_no_column_is_listed_twice():
    assert len(csv_export.COLUMNS) == len(set(csv_export.COLUMNS))


# ------------------------------------------------------------------ the workbook
def test_the_rank_sheet_carries_the_cross_sector_columns_populated():
    ws = _book(_universe())["Rank"]
    hdr = next(i for i, row in enumerate(ws.iter_rows(values_only=True), start=1)
               if row and row[0] == "ticker")
    cols = {c.value: c.column for c in ws[hdr]}
    assert "sector_rank" in cols and "sector_neutral_score" in cols
    vals = [ws.cell(row=r, column=cols["sector_rank"]).value
            for r in range(hdr + 1, hdr + 1 + 16)]
    assert all(isinstance(v, int) for v in vals), "the headline column rendered blank"


def test_a_workbook_built_from_rows_that_never_saw_neutralize_still_fills_the_column():
    """build_workbook must not depend on the crawl having run first."""
    ws = _book(_universe())["Rank"]          # _universe() is deliberately NOT neutralized
    hdr = next(i for i, row in enumerate(ws.iter_rows(values_only=True), start=1)
               if row and row[0] == "ticker")
    cols = {c.value: c.column for c in ws[hdr]}
    assert ws.cell(row=hdr + 1, column=cols["sector_neutral_score"]).value is not None


def test_the_rank_sheet_says_which_column_answers_the_cross_sector_question():
    txt = _text(_book(_universe())["Rank"])
    assert "sector_neutral_score" in txt
    assert "composite_strict, NOT over the score column" in txt, \
        "the percentile's basis must be stated - score and composite_strict disagree"


def test_the_note_quotes_a_spread_measured_from_the_rows_not_from_memory():
    txt = _text(_book(_universe())["Rank"])
    # IT median 59, Real Estate median 32 -> 27 points, computed live
    assert "27-point spread" in txt
    assert "Information Technology 59" in txt and "Real Estate 32" in txt


def test_a_null_score_does_not_scramble_the_sector_median():
    """NaN is a float and every comparison against it is False, so it does not skew a
    sorted list - it silently reorders it. Caught on the real 1,501-row parquet: three
    null scores made the Real Estate median read 12, its MINIMUM, instead of 33."""
    rows = _universe()
    rows.append({"ticker": "NULL", "name": "No score", "cik": "99",
                 "sector": "Real Estate", "sub_industry": "REITs",
                 "composite_strict": 30, "score": float("nan"),
                 "band": "INSUFFICIENT_DATA", "coverage": 0.4})
    med = dict((s, m) for s, _n, m, _b in xlsx_export._sector_medians(rows))
    assert med["Real Estate"] == 32, med


def test_the_leaderboard_does_not_rank_a_null_score_first():
    rows = _universe()
    rows.append({"ticker": "NULL", "name": "No score", "cik": "99",
                 "sector": "Real Estate", "sub_industry": "REITs",
                 "composite_strict": 30, "score": float("nan"),
                 "band": "INSUFFICIENT_DATA", "coverage": 0.4})
    ws = _book(rows)["Sectors"]
    hdr = next(i for i, row in enumerate(ws.iter_rows(values_only=True), start=1)
               if row and row[1] == "rank_in_sector")
    tickers = [ws.cell(row=r, column=4).value for r in range(hdr + 1, hdr + 20)]
    assert "NULL" not in tickers


def test_the_note_survives_a_universe_too_thin_to_measure_a_spread():
    rows = [{"ticker": "A", "sector": "S", "sub_industry": "S",
             "composite_strict": 50, "score": 50, "coverage": 0.9}]
    txt = _text(_book(rows)["Rank"])
    assert "COMPARING ACROSS SECTORS" in txt


def test_the_sectors_sheet_ranks_the_top_names_within_each_sector():
    ws = _book(_universe())["Sectors"]
    txt = _text(ws)
    assert "Top 10 within each sector" in txt
    hdr = next(i for i, row in enumerate(ws.iter_rows(values_only=True), start=1)
               if row and row[1] == "rank_in_sector")
    lead = [ws.cell(row=hdr + 1, column=c).value for c in (1, 2, 3, 4)]
    assert lead == ["Information Technology", 1, 8, "IT7"]


def test_the_leaderboard_shows_the_neutral_score_beside_the_raw_one():
    ws = _book(_universe())["Sectors"]
    hdr = next(i for i, row in enumerate(ws.iter_rows(values_only=True), start=1)
               if row and row[1] == "rank_in_sector")
    header = [c.value for c in ws[hdr]]
    assert "score" in header and "sector_neutral_score" in header


def test_findings_answers_the_cross_sector_question_and_claims_nothing():
    txt = _text(_book(_universe())["Findings"])
    assert "USE sector_neutral_score, NOT score" in txt
    assert "PRESENTATION change" in txt
    assert "no claim that either column predicts returns" in txt


def test_findings_does_not_present_the_change_as_a_result():
    """A percentile removes sector spread BY CONSTRUCTION - E05's F1 is the warning.

    Scanned on the sector-neutral ROW only. The sheet as a whole says "validated" and
    "better" all over the place about other studies, and a whole-sheet substring scan
    would pass or fail for reasons that have nothing to do with this change.
    """
    ws = _book(_universe())["Findings"]
    detail = next(row[2] for row in ws.iter_rows(values_only=True)
                  if row and row[1] == "USE sector_neutral_score, NOT score")
    assert "nothing was rescored, no weight moved" in detail
    for claim in ("improves the ranking", "ranks better", "beats"):
        assert claim not in detail

