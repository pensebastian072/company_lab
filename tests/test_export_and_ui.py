"""Exports and the dashboard.

The UI tests all assert HTTP 200 with missing or corrupt artifacts. A dashboard
that 500s because a crawl is mid-write is a dashboard you stop trusting.
"""
from __future__ import annotations

import io
import json
import sys

import pytest

from clab import config
from clab.export import csv_export, xlsx_export

ROW = {
    "ticker": "TEST", "cik": "0000000001", "name": "Test Co",
    "sector": "Information Technology", "sub_industry": "Software",
    "profile": "STANDARD", "as_of": "2026-08-08T00:00:00+00:00",
    "data_through": "2026-06-30", "price_as_of": "2026-08-07",
    "composite_strict": 72, "composite_normalized": 80.0, "composite_ex_entry": 69,
    "band": "INVESTABLE", "coverage": 0.9, "quant_only_50": 38,
    "quant_normalized": 84.4, "quant_band": "HIGH_CONVICTION", "quant_coverage": 0.9,
    "qual_available": True, "points_earned": 72, "points_available": 90,
    "n_subtests_no_data": 2, "fe": 13, "bs": 9, "va": 8, "er": 4, "en": 4,
    "sg": 16, "bq": 11, "mg": 7, "roic": 0.21, "operating_margin": 0.28,
    "revenue_ttm": 1.23e10, "net_debt_ebitda": 0.4, "pe": 24.5,
    "stress_verdict": "SURVIVES_COMFORTABLY",
}

CARD = {
    "ticker": "TEST", "cik": "0000000001", "name": "Test Co",
    "metrics": {"price": 100.0, "pe": 24.5, "market_cap": 1e11},
    "components": {
        "FE": {"label": "Financial Engine", "max_points": 15, "earned_points": 13,
               "available_points": 15, "applicable_points": 15,
               "not_applicable_points": 0, "no_data_points": 0, "coverage": 1.0,
               "is_llm": False, "note": "",
               "subtests": [{"key": "fe_roic", "label": "ROIC", "max_points": 2,
                             "earned": 2, "effective": 2, "status": "scored",
                             "inputs": {"roic": 0.21},
                             "provenance": {"source": "edgar",
                                            "tags_used": {"roic": "derived"}},
                             "rationale": "", "evidence": None,
                             "evidence_unverified": False, "override": None,
                             "threshold_note": "ROIC 21.0%"}]},
        "BS": {"label": "Balance Sheet", "max_points": 10, "earned_points": 9,
               "available_points": 10, "applicable_points": 10,
               "not_applicable_points": 0, "no_data_points": 0, "coverage": 1.0,
               "is_llm": False, "note": "",
               "subtests": [{"key": "bs_stress_test", "label": "Downturn survival test",
                             "max_points": 2, "earned": 2, "effective": 2,
                             "status": "scored",
                             "inputs": {"verdict": "SURVIVES_COMFORTABLY",
                                        "points": 2,
                                        "inputs": {"revenue_ttm": 1.0,
                                                   "operating_margin": 0.3},
                                        "simulated": {"revenue": 0.8, "fcf": 0.1}},
                             "provenance": {}, "rationale": "", "evidence": None,
                             "evidence_unverified": False, "override": None,
                             "threshold_note": "-> SURVIVES_COMFORTABLY"}]},
        "VA": {"label": "Valuation", "max_points": 15, "earned_points": 8,
               "available_points": 12, "applicable_points": 15,
               "not_applicable_points": 0, "no_data_points": 3, "coverage": 0.8,
               "is_llm": False, "note": "",
               "subtests": [{"key": "va_reverse_dcf",
                             "label": "Reverse DCF", "max_points": 5, "earned": 3,
                             "effective": 3, "status": "scored",
                             "inputs": {"implied_growth": 0.08, "wacc": 0.09,
                                        "sensitivity": {
                                            "wacc_minus": {"implied_growth": 0.05},
                                            "wacc_plus": {"implied_growth": 0.11}},
                                        "reference_growth": 0.1,
                                        "gap_vs_reference": -0.02},
                             "provenance": {"source": "derived"}, "rationale": "",
                             "evidence": None, "evidence_unverified": False,
                             "override": None, "threshold_note": "requires 8%"}]},
        "SG": {"label": "Structural Growth", "max_points": 20, "earned_points": 16,
               "available_points": 20, "applicable_points": 20,
               "not_applicable_points": 0, "no_data_points": 0, "coverage": 1.0,
               "is_llm": True, "note": "",
               "subtests": [{"key": "sg_tam_expanding", "label": "TAM expanding",
                             "max_points": 4, "earned": 3, "effective": 3,
                             "status": "scored", "inputs": {},
                             "provenance": {"source": "llm", "model": "qwen2.5:7b"},
                             "rationale": "the model's sentence",
                             "evidence": "a verbatim quote",
                             "evidence_unverified": False, "override": None,
                             "threshold_note": "model score 3/4"}]},
    },
}


# ------------------------------------------------------------------ CSV
def test_csv_has_stable_header_and_bom():
    data = csv_export.to_bytes([ROW])
    assert data.startswith(b"\xef\xbb\xbf")            # Excel-friendly BOM
    text = data.decode("utf-8-sig")
    header = text.splitlines()[0].split(",")
    assert header == list(csv_export.COLUMNS)
    assert header[0] == "ticker"
    assert "quant_only_50" in header and "quant_band" in header


def test_csv_uses_crlf():
    assert "\r\n" in csv_export.to_string([ROW])


def test_csv_tolerates_missing_keys():
    text = csv_export.to_string([{"ticker": "X"}])
    assert text.splitlines()[1].startswith("X,")


def test_scorecard_long_rows_are_the_audit_trail():
    rows = csv_export.scorecard_long_rows(CARD)
    keys = {(r["component"], r["subtest"]) for r in rows}
    assert ("FE", "fe_roic") in keys
    assert ("SG", "sg_tam_expanding") in keys
    llm = [r for r in rows if r["source"] == "llm"][0]
    assert llm["evidence"] == "a verbatim quote"
    assert llm["rationale"] == "the model's sentence"


# ------------------------------------------------------------------ XLSX
def test_workbook_has_all_sheets():
    from openpyxl import load_workbook

    data = xlsx_export.to_bytes([ROW], cards=[CARD])
    wb = load_workbook(io.BytesIO(data))
    # TAB ORDER IS A HUMAN CONTRACT and it changed 2026-09-07: the two RANKING sheets
    # now lead, and everything that exists for a machine or an auditor moved right.
    # Rank then Sectors, because those are what the workbook is opened to see.
    # Changes/Findings still sit near the front: a workbook showing confident-looking
    # scores with no measured predictive power beside them is misleading by omission.
    # DCF sits behind Valuation: the reverse DCF answers "what growth does the price
    # require", the forward one answers "what is a share worth", and reading either
    # without the other is how a fair value gets mistaken for a target.
    # Components/Subtests/Claims/Meta are back-end dumps and provenance, so they are
    # last. The full order lives in xlsx_export.SHEET_ORDER; this pins what a build
    # with no external store actually produces.
    # The external sheets appear only when an external store exists on the box, so the
    # core order is asserted with them removed, and their POSITION is asserted
    # separately. Pinning a 15-name list here would make this test pass or fail on
    # whether a database happens to be present, which is not what it is checking.
    external = {"External", "Industries", "Claims"}
    core = [n for n in wb.sheetnames if n not in external]
    assert core == ["Rank", "Sectors", "Changes", "Findings", "Entry",
                    "Valuation", "DCF", "Thesis", "Stress",
                    "Components", "Subtests", "Meta"]
    for name in external & set(wb.sheetnames):
        assert wb.sheetnames.index(name) > wb.sheetnames.index("Stress")
        assert wb.sheetnames.index(name) < wb.sheetnames.index("Meta")


def test_entry_sheet_carries_the_entry_columns():
    from openpyxl import load_workbook

    row = dict(ROW, pe_current=22.5, pe_pctile_own=0.12, pe_median_own=40.0,
               pct_vs_ema20_monthly=-0.02, en_confluence=True)
    wb = load_workbook(io.BytesIO(xlsx_export.to_bytes([row], cards=[CARD])))
    ws = wb["Entry"]
    header = [c.value for c in ws[1]]
    for col in ("pe_pctile_own", "pct_vs_ema20_monthly", "pct_vs_ema72_weekly",
                "en_confluence", "en"):
        assert col in header
    assert ws.cell(row=2, column=header.index("pe_pctile_own") + 1).value == 0.12
    assert ws.auto_filter.ref is not None


def test_meta_sheet_leads_with_the_disclaimer():
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx_export.to_bytes([ROW], cards=[CARD])))
    a1 = str(wb["Meta"]["A1"].value or "")
    assert "ADVISORY" in a1 and "SHADOW" in a1
    assert "LLM" in a1


def test_meta_sheet_states_the_horizon_and_keeps_its_key_value_header():
    """F4. The horizon line was declared in the rubric for a day and used nowhere.

    Also pins the header row, which the horizon line pushed down: bolding a fixed row 3
    would have silently bolded a blank row instead.
    """
    from openpyxl import load_workbook

    from clab.scoring import rubric

    wb = load_workbook(io.BytesIO(xlsx_export.to_bytes([ROW], cards=[CARD])))
    ws = wb["Meta"]
    assert rubric.HORIZON_LABEL in str(ws["A2"].value or "")

    header = next((r for r in range(1, 12)
                   if str(ws.cell(row=r, column=1).value or "") == "key"), None)
    assert header is not None, "the key/value header vanished"
    assert ws.cell(row=header, column=1).font.bold, "header row is no longer bold"
    keys = {str(ws.cell(row=r, column=1).value): ws.cell(row=r, column=2).value
            for r in range(header + 1, ws.max_row + 1)}
    assert keys.get("default_horizon") == rubric.DEFAULT_HORIZON


def test_horizon_is_actually_surfaced_by_the_ui(client):
    """A rubric constant nobody renders is not an implemented fix."""
    from clab.scoring import rubric

    body = client.get("/").get_data(as_text=True)
    assert rubric.HORIZON_LABEL in body
    assert rubric.HORIZON_NOTE[:40] in body


def _rank_header_row(ws) -> int:
    """The Rank sheet carries a partial-rollout note above the header while the LLM
    half is still being filled in, so the header row is found, never assumed."""
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row and row[0] == "ticker":
            return i
    raise AssertionError("no header row on the Rank sheet")


def test_rank_sheet_is_filterable_and_frozen():
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx_export.to_bytes([ROW], cards=[CARD])))
    ws = wb["Rank"]
    hdr = _rank_header_row(ws)
    assert ws.freeze_panes == f"B{hdr + 1}"
    assert ws.auto_filter.ref is not None
    assert ws.auto_filter.ref.startswith(f"A{hdr}:")
    assert ws.conditional_formatting is not None


def test_ticker_cell_links_to_the_local_scorecard():
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx_export.to_bytes([ROW], cards=[CARD])))
    ws = wb["Rank"]
    val = str(ws.cell(row=_rank_header_row(ws) + 1, column=1).value)
    assert val.startswith("=HYPERLINK(")
    assert f"{config.UI_HOST}:{config.UI_PORT}/c/TEST" in val


def test_subset_of_sheets_can_be_requested():
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(
        xlsx_export.to_bytes([ROW], cards=[CARD], sheets=("rank", "meta"))))
    assert wb.sheetnames == ["Rank", "Meta"]


def test_workbook_survives_an_empty_universe():
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx_export.to_bytes([], cards=[])))
    assert "Meta" in wb.sheetnames


# ------------------------------------------------------------------ UI
@pytest.fixture()
def client():
    from clab.ui.app import app

    app.config.update(TESTING=True)
    return app.test_client()


def test_ui_never_imports_the_llm_scorer():
    """The LLM must not sit on any request path. Same discipline as
    hq-trading-system's receiver not importing llm_macro_analyst.

    Checked in a FRESH interpreter: within one pytest session another test module
    imports clab.qual, so an in-process sys.modules assertion would only be testing
    test ordering. A subprocess is the real check.
    """
    import subprocess

    code = (
        "import sys; import clab.ui.app; "
        "bad = [m for m in sys.modules if m.startswith('clab.qual')]; "
        "print('LEAKED:' + ','.join(bad) if bad else 'CLEAN')"
    )
    r = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True,
                       cwd=str(config.BASE_DIR), timeout=120)
    assert r.returncode == 0, r.stderr
    assert "CLEAN" in r.stdout, r.stdout


def test_ui_binds_loopback_only():
    assert config.UI_HOST == "127.0.0.1"
    src = (config.BASE_DIR / "clab" / "ui" / "app.py").read_text(encoding="utf-8")
    assert "0.0.0.0" not in src


@pytest.mark.parametrize("path", ["/", "/health", "/c/AAPL", "/c/NOTAREALTICKER"])
def test_pages_render(client, path):
    r = client.get(path)
    assert r.status_code == 200


@pytest.mark.parametrize("path", [
    "/api/state", "/api/health", "/api/company/AAPL", "/api/thesis/AAPL",
    "/api/company/!!bad!!", "/api/refresh",
])
def test_api_always_200(client, path):
    r = client.get(path)
    assert r.status_code == 200
    assert r.get_json() is not None


def test_state_payload_shape(client):
    d = client.get("/api/state").get_json()
    for key in ("ok", "advisory", "rows", "n", "flag", "components"):
        assert key in d
    assert len(d["components"]) == 8
    llm = [c for c in d["components"] if c["is_llm"]]
    # V3 P0: MG dropped out. Every MG sub-test is source=edgar and every MG payload
    # carries ceo={} - the model contributes only a display-only summary string.
    assert {c["code"] for c in llm} == {"SG", "BQ"}


def test_health_payload_shape(client):
    d = client.get("/api/health").get_json()
    for key in ("ok", "advisory", "flag", "stale", "manifest", "freshness"):
        assert key in d


def test_corrupt_scores_table_degrades_to_200(client, monkeypatch, tmp_path):
    bad = tmp_path / "scores.parquet"
    bad.write_bytes(b"this is not parquet")
    monkeypatch.setattr(config, "SCORES_PARQUET", bad)
    from clab.ui import state as state_mod

    monkeypatch.setattr(state_mod.config, "SCORES_PARQUET", bad)
    st = state_mod.build_state()
    assert st["ok"] is False
    assert st["rows"] == []
    assert "could not read" in st["reason"]


def test_missing_scores_table_degrades_to_200(monkeypatch, tmp_path):
    from clab.ui import state as state_mod

    monkeypatch.setattr(state_mod.config, "SCORES_PARQUET", tmp_path / "nope.parquet")
    st = state_mod.build_state()
    assert st["ok"] is False
    assert "does not exist" in st["reason"]


def test_export_routes_return_attachments(client):
    r = client.get("/export/scores.csv")
    assert r.status_code == 200
    assert "attachment" in r.headers["Content-Disposition"]
    r2 = client.get("/export/company_lab.xlsx?sheets=rank,meta")
    assert r2.status_code == 200
    assert "attachment" in r2.headers["Content-Disposition"]


def test_bad_ticker_export_does_not_500(client):
    r = client.get("/export/company/!!!.csv")
    assert r.status_code == 200


def test_flag_contract_keys_present_when_written(tmp_path, monkeypatch):
    """Sibling repos read this file; the key contract must not drift."""
    from clab.runner import batch

    out = tmp_path / "state.json"
    monkeypatch.setattr(batch.config, "STATE_FLAG", out)
    batch.write_state_flag({"n_ok": 3, "n_failed": 0, "tier": "sp500",
                            "batches_done": 1, "batches_total": 1}, [ROW])
    d = json.loads(out.read_text(encoding="utf-8"))
    for key in ("as_of", "data_through", "stale", "status", "promoted",
                "symbols_scored", "coverage_median", "qual_available"):
        assert key in d
    assert d["promoted"] is False        # nothing here is ever promoted
