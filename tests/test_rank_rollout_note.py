"""The Rank sheet's partial-rollout note, and the row offsets it shifts.

Why the note exists: the judgement half is 50 of the 100 points and is being filled in
at 100 companies a day. Measured mid-rollout, median composite_strict was 33 for
companies that had it and 10 for those that did not - so an unscored company CANNOT
appear near the top, and the ordering reflects who has been scored as much as who is
good. Ranking a half-scored universe by raw points without saying so is misleading.

Why the offsets are tested: adding a row above a header is exactly the off-by-one that
already bolded a blank row and aimed a chart one row high elsewhere in this module. The
header, autofilter, colour scale and number formats must follow the header row, not a
hardcoded 1.

Since 2026-08-25 there is a SECOND row above the header - the cross-sector note, which is
always present - so the header sits at 2 with no rollout note and 3 with one. That is
exactly the shift these tests exist to catch.
"""
from __future__ import annotations

import io

import pytest

from clab.export import xlsx_export


def _row(t, *, qual: bool, score: int = 50):
    r = {"ticker": t, "cik": f"{hash(t) % 10**10:010d}", "name": f"{t} Inc",
         "sector": "Information Technology", "composite_strict": score,
         "band": "WATCHLIST", "coverage": 0.9}
    for c in ("sg", "bq", "mg"):
        r[f"{c}_available"] = 15 if qual else 0
    return r


def _book(rows):
    from openpyxl import load_workbook

    data = xlsx_export.to_bytes(rows, cards=[], sheets=("rank",))
    return load_workbook(io.BytesIO(data))["Rank"]


def _header_row(ws):
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row and row[0] == "ticker":
            return i
    raise AssertionError("no header row found")


PARTIAL = [_row("AAA", qual=True, score=80), _row("BBB", qual=False, score=20),
           _row("CCC", qual=False, score=15)]
COMPLETE = [_row("AAA", qual=True, score=80), _row("BBB", qual=True, score=60)]


def _abstained(t, *, score: int = 20):
    """Been through the model, which answered on SG and declined on BQ.

    SG is the tell: it has no non-LLM points, so SG available means the call was made.
    """
    r = _row(t, qual=False, score=score)
    r["sg_available"] = 15
    r["bq_available"] = 0
    return r


ABSTAINED = [_row("DDD", qual=True, score=80), _abstained("EEE"), _abstained("FFF")]


def test_note_appears_while_the_rollout_is_partial():
    ws = _book(PARTIAL)
    assert str(ws["A1"].value or "").startswith("PARTIAL")
    # counts the companies STILL TO DO, which is what a wait resolves
    assert "2 of 3" in ws["A1"].value
    assert _header_row(ws) == 3


def test_abstention_is_not_reported_as_an_unfinished_crawl():
    """The two reasons BQ can be missing must never be added together.

    Measured 2026-08-18: 1,035 of 1,498 companies had BQ. The banner read "only 1,035 of
    1,498 ... (100 added per day)", which promised a rollout that would finish. It never
    would: exactly ONE company was unscored and the other 462 had been through the model,
    which declined to answer (E12). A reader waiting for that banner to clear waits
    forever, so the crawl-progress claim and the abstention finding are now separate.
    """
    ws = _book(ABSTAINED)
    note = str(ws["A1"].value or "")
    assert "PARTIAL" not in note, note
    assert "COMPLETE" in note
    assert "2 of 3" in note                     # the two the model declined
    assert "E12" in note
    assert "declined" in note.lower()
    assert _header_row(ws) == 3                 # rollout note row 1, cross-sector row 2


def test_abstention_alone_leaves_the_crawl_complete():
    roll = xlsx_export._qual_rollout(ABSTAINED)
    assert roll["complete"] is True, "the crawl IS done; only the model refused"
    assert roll["not_run"] == 0
    assert roll["abstained"] == 2
    assert roll["with_qual"] == 1


def test_the_two_reasons_are_counted_separately():
    roll = xlsx_export._qual_rollout(PARTIAL + ABSTAINED)
    assert roll["not_run"] == 2                  # PARTIAL's BBB and CCC
    assert roll["abstained"] == 2                # ABSTAINED's EEE and FFF
    assert roll["without_qual"] == roll["not_run"] + roll["abstained"]
    assert roll["complete"] is False             # an unreached company outranks the rest


def test_no_note_once_every_company_has_the_judgement_half():
    ws = _book(COMPLETE)
    assert _header_row(ws) == 2   # the cross-sector note is unconditional and keeps row 1
    assert "PARTIAL" not in str(ws["A1"].value or "")


@pytest.mark.parametrize("rows,want_hdr", [(PARTIAL, 3), (COMPLETE, 2)])
def test_header_is_styled_on_the_row_it_was_written_to(rows, want_hdr):
    ws = _book(rows)
    hdr = _header_row(ws)
    assert hdr == want_hdr
    assert ws.cell(row=hdr, column=1).font.bold, "header not bold"
    assert ws.freeze_panes == f"B{hdr + 1}"


@pytest.mark.parametrize("rows", [PARTIAL, COMPLETE])
def test_autofilter_starts_at_the_header_and_covers_every_company(rows):
    ws = _book(rows)
    hdr = _header_row(ws)
    ref = ws.auto_filter.ref
    assert ref.startswith(f"A{hdr}:"), ref
    assert ref.endswith(str(hdr + len(rows))), ref


@pytest.mark.parametrize("rows", [PARTIAL, COMPLETE])
def test_colour_scale_covers_the_data_rows_only(rows):
    ws = _book(rows)
    hdr = _header_row(ws)
    ranges = [str(r.sqref) for r in ws.conditional_formatting]
    assert ranges, "the colour scale on the headline score is missing"
    assert any(f"{hdr + 1}" in r for r in ranges), ranges


def test_every_company_is_present_under_the_note():
    ws = _book(PARTIAL)
    hdr = _header_row(ws)
    seen = {ws.cell(row=r, column=1).value for r in range(hdr + 1, ws.max_row + 1)}
    # tickers are rendered as HYPERLINK formulas, so match on the contained text
    joined = " ".join(str(s) for s in seen)
    for t in ("AAA", "BBB", "CCC"):
        assert t in joined


def test_mg_alone_does_not_count_as_llm_scored():
    """MG is deliberately a MIXTURE: 9.4 of its 15 points are measured from filings, so
    it has available points with no LLM at all. Counting it reported 1,481 of 1,497
    companies as LLM-scored when the true figure was 698."""
    r = _row("MMM", qual=False)
    r["mg_available"] = 9          # the measured part of MG, no LLM involved
    roll = xlsx_export._qual_rollout([r])
    assert roll["with_qual"] == 0
    assert roll["without_qual"] == 1


def test_partial_llm_half_does_not_count():
    r = _row("PPP", qual=False)
    r["sg_available"] = 20         # SG present, BQ missing - still a hole
    roll = xlsx_export._qual_rollout([r])
    assert roll["with_qual"] == 0


def test_rollout_helper_counts_by_real_component_availability():
    roll = xlsx_export._qual_rollout(PARTIAL)
    assert roll["with_qual"] == 1
    assert roll["without_qual"] == 2
    assert roll["complete"] is False
    assert roll["median_with"] == 80
    assert xlsx_export._qual_rollout(COMPLETE)["complete"] is True


def test_findings_sheet_states_the_rollout_first():
    from openpyxl import load_workbook

    data = xlsx_export.to_bytes(PARTIAL, cards=[], sheets=("findings",))
    ws = load_workbook(io.BytesIO(data))["Findings"]
    text = "\n".join(str(c) for row in ws.iter_rows(values_only=True)
                     for c in row if c is not None)
    assert "Is the whole universe scored on all 100 points?" in text
    assert "cannot appear near the top" in text.lower() or "CANNOT appear" in text


def test_both_reasons_at_once_name_the_bigger_one_first():
    """The real 2026-08-18 shape: 1 company unreached, 462 refused by the model.

    A first pass branched on "is the crawl complete" and so reported only the 1, burying
    the 462 behind a rounding error. Both must appear, larger first.
    """
    roll = {"total": 1498, "not_run": 1, "abstained": 462}
    note = xlsx_export._rollout_banner(roll)
    assert "462 of 1498" in note
    assert "1 of 1498" in note
    assert note.index("462") < note.index("1 of 1498"), "the bigger reason must lead"
    assert "E12" in note


def test_the_banner_disappears_only_when_nothing_is_missing():
    assert xlsx_export._rollout_banner({"total": 10, "not_run": 0, "abstained": 0}) is None
    assert xlsx_export._rollout_banner({"total": 0}) is None
