"""Two-level sector reporting: 11 GICS sectors on top, sub-industries beneath.

The user's point, and it is the right one: 124 "sectors" is unreadable, but the level
below the eleven is where the signal is - which sub-industries inside a sector lead and
lag says more about where the market is going than the sector average does.

So the SIC description is not discarded, it is moved to the level it actually belongs
to. It was previously sitting in `sector` (124 pseudo-sectors) and these companies had
an EMPTY sub_industry, which meant no peer group at all and every peer-relative
valuation sub-test silently unavailable for them.
"""
from __future__ import annotations

import io

from clab.export import xlsx_export


def _row(t, sector, sub, score, band="WATCHLIST"):
    return {"ticker": t, "cik": f"{abs(hash(t)) % 10**10:010d}", "name": f"{t} Inc",
            "sector": sector, "sub_industry": sub, "composite_strict": score,
            "band": band, "coverage": 0.95,
            "sg_available": 20, "bq_available": 15, "mg_available": 15}


def _rows():
    out = []
    # Software: 5 companies, high scores
    for i, s in enumerate((80, 78, 76, 74, 72)):
        out.append(_row(f"SW{i}", "Information Technology", "Application Software", s))
    # Semis: 5 companies, lower
    for i, s in enumerate((60, 58, 56, 54, 52)):
        out.append(_row(f"SC{i}", "Information Technology", "Semiconductors", s))
    # a thin group that must not be ranked
    out.append(_row("TINY", "Information Technology", "Quantum Widgets", 99))
    # another sector
    for i, s in enumerate((40, 38, 36, 34, 32)):
        out.append(_row(f"UT{i}", "Utilities", "Electric Utilities", s))
    return out


def _sectors_sheet(rows=None):
    from openpyxl import load_workbook

    data = xlsx_export.to_bytes(rows if rows is not None else _rows(), cards=[],
                                sheets=("sectors",))
    return load_workbook(io.BytesIO(data))["Sectors"]


def _text(ws):
    return "\n".join(str(c) for row in ws.iter_rows(values_only=True)
                     for c in row if c is not None)


def _sub_rows(ws):
    """Rows of the sub-industry table, as (sector, sub, n, mean, vs_sector)."""
    out, seen_header = [], False
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "sector" and row[1] == "sub_industry":
            seen_header = True
            continue
        if seen_header and row and row[0] and row[1] and isinstance(row[2], int):
            out.append(row[:5])
    return out


def test_sub_industry_table_exists_and_is_ranked_within_sector():
    ws = _sectors_sheet()
    subs = _sub_rows(ws)
    it = [r for r in subs if r[0] == "Information Technology"]
    assert [r[1] for r in it] == ["Application Software", "Semiconductors"], (
        "sub-industries must be ranked by mean score within their sector")


def test_vs_sector_column_shows_the_gap_to_the_sector_average():
    ws = _sectors_sheet()
    subs = {r[1]: r for r in _sub_rows(ws)}
    # The sector average is over EVERY company in the sector, including those in
    # sub-industries too thin to rank: TINY (99) counts here even though "Quantum
    # Widgets" is not listed as a sub-industry. IT mean =
    # (80+78+76+74+72+60+58+56+54+52+99)/11 = 69.
    assert round(subs["Application Software"][3], 1) == 76.0
    assert round(subs["Application Software"][4], 1) == 7.0
    assert round(subs["Semiconductors"][4], 1) == -13.0


def test_thin_groups_are_listed_but_not_ranked():
    """A two-company 'sub-industry' is a company name, not an industry."""
    ws = _sectors_sheet()
    subs = [r[1] for r in _sub_rows(ws)]
    assert "Quantum Widgets" not in subs
    txt = _text(ws)
    assert "Quantum Widgets (1)" in txt
    assert f"fewer than {xlsx_export.MIN_SUB_INDUSTRY_N} companies" in txt


def test_thin_group_does_not_distort_its_sector_ranking():
    """TINY scores 99 but must not make Quantum Widgets the top IT sub-industry."""
    ws = _sectors_sheet()
    it = [r for r in _sub_rows(ws) if r[0] == "Information Technology"]
    assert it[0][1] == "Application Software"


def test_sector_level_still_present_above_the_sub_industry_block():
    ws = _sectors_sheet()
    txt = _text(ws)
    assert "SECTOR BET" in txt, "the sector-level finding must stay"
    assert "spread between the highest and lowest sector" in txt
    assert "Sub-industries" in txt


def test_missing_sub_industry_is_labelled_not_dropped():
    rows = _rows()
    for r in rows:
        r["sub_industry"] = ""
    txt = _text(_sectors_sheet(rows))
    assert "Unclassified" in txt


def test_empty_universe_does_not_break_the_sheet():
    ws = _sectors_sheet([])
    assert ws.max_row >= 1


def test_sic_descriptions_merge_into_their_gics_sub_industry():
    """Leaving them raw split real peer groups: "Semiconductors" (26) beside
    "Semiconductors & Related Devices" (9), "Application Software" (33) beside
    "Services-Prepackaged Software" (9)."""
    from clab.research.sector_fill import normalize_sub_industry as norm

    assert norm("Semiconductors & Related Devices") == "Semiconductors"
    assert norm("Services-Prepackaged Software") == "Application Software"
    assert norm("Pharmaceutical Preparations") == "Pharmaceuticals"
    assert norm("State Commercial Banks") == "Regional Banks"
    assert norm("National Commercial Banks") == "Diversified Banks"
    assert norm("Crude Petroleum & Natural Gas") == "Oil & Gas Exploration & Production"


def test_a_gics_label_passes_through_untouched():
    from clab.research.sector_fill import normalize_sub_industry as norm

    for label in ("Semiconductors", "Application Software", "Regional Banks",
                  "Health Care Equipment"):
        assert norm(label) == label


def test_an_unmapped_sic_description_is_kept_not_guessed():
    """A truthful odd label beats a confident wrong merge."""
    from clab.research.sector_fill import normalize_sub_industry as norm

    assert norm("Ball & Roller Bearings") != ""
    assert norm("Some Entirely Novel Industry") == "Some Entirely Novel Industry"
    assert norm("") == ""
    assert norm(None) == ""


def test_normalisation_is_insensitive_to_case_and_spacing():
    """The SEC's descriptions contain double spaces, e.g. 'Retail-Eating  Places'."""
    from clab.research.sector_fill import normalize_sub_industry as norm

    assert norm("Retail-Eating  Places") == "Restaurants"
    assert norm("  services-prepackaged SOFTWARE  ") == "Application Software"


def test_every_mapping_target_is_a_plausible_gics_name():
    """Guards against a typo silently creating a new one-company group."""
    from clab.research.sector_fill import SIC_DESC_TO_GICS_SUB as M

    for key, target in M.items():
        assert key == key.lower().strip(), f"key {key!r} must be lowercased"
        assert target and target[0].isupper(), f"{target!r} does not look like a label"
        assert target not in M, f"{target!r} is both a key and a target"


def test_engine_puts_the_sic_description_at_sub_industry_level():
    """Guards the regression: sic_desc must be the sub-industry fallback, not sector."""
    import inspect

    from clab.runner import engine

    src = inspect.getsource(engine)
    # sic_desc feeds the SUB-INDUSTRY fallback (via normalize_sub_industry), never the
    # sector, and the sector comes from the SIC code mapping instead.
    assert "normalize_sub_industry(sub_industry or (sic_desc" in src
    assert "sector=resolved_sector" in src
    assert 'sector=sector or (sic_desc or "")' not in src
