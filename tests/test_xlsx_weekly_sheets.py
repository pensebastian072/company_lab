"""The three sheets that make the workbook a weekly read: Changes, Sectors, Findings.

What is worth pinning here is honesty, not formatting:

  * the Changes sheet must never present a two-day gap as a week;
  * the Findings sheet must carry the FAILURES, because a workbook showing 500
    confident-looking scores with no measured predictive power beside them is
    misleading by omission;
  * the Sectors sheet must show the spread that E03 R1 found, since two companies in
    different sectors are partly being compared on their sectors.
"""
from __future__ import annotations

import datetime as dt
import io

import pandas as pd
import pytest

from clab import config
from clab.export import xlsx_export
from clab.scoring import rubric


def _rows():
    return [
        {"ticker": "AAA", "cik": "1", "sector": "Information Technology",
         "composite_strict": 82, "band": "HIGH_CONVICTION", "coverage": 0.95},
        {"ticker": "BBB", "cik": "2", "sector": "Information Technology",
         "composite_strict": 60, "band": "WATCHLIST", "coverage": 0.90},
        {"ticker": "CCC", "cik": "3", "sector": "Utilities",
         "composite_strict": 41, "band": "REJECT", "coverage": 0.85},
        {"ticker": "DDD", "cik": "4", "sector": "Utilities",
         "composite_strict": 39, "band": "REJECT", "coverage": 0.80},
    ]


def _book(rows=None, sheets=("sectors", "findings", "changes")):
    from openpyxl import load_workbook

    data = xlsx_export.to_bytes(rows if rows is not None else _rows(), cards=[],
                                sheets=sheets)
    return load_workbook(io.BytesIO(data))


def _cells(ws):
    return [c for row in ws.iter_rows(values_only=True) for c in row if c is not None]


def _text(ws):
    return "\n".join(str(c) for c in _cells(ws))


# ------------------------------------------------------------------ Sectors
def test_sectors_sheet_ranks_sectors_and_states_the_spread():
    ws = _book()["Sectors"]
    txt = _text(ws)
    assert "SECTOR BET" in txt, "the finding that motivates the sheet is missing"
    assert "spread between the highest and lowest sector" in txt
    # IT mean 71, Utilities mean 40 -> 31 points
    assert "31.0 points" in txt
    hdr = next(i for i, row in enumerate(ws.iter_rows(values_only=True), start=1)
               if row and row[0] == "sector")
    order = [ws.cell(row=r, column=1).value
             for r in range(hdr + 1, hdr + 3)]
    assert order == ["Information Technology", "Utilities"], "sectors not sorted by score"


def test_sectors_sheet_carries_a_native_chart():
    ws = _book()["Sectors"]
    assert len(ws._charts) == 1, "the bar chart is the point of the sheet"


@pytest.mark.parametrize("sheet,first_header", [("Sectors", "sector"),
                                               ("Findings", "question")])
def test_header_row_is_bold_on_the_row_it_is_actually_written_to(sheet, first_header):
    """`ws.append([])` advances the write cursor but writes no cells, so max_row does
    NOT move - predicting the header row instead of reading it back is off by one and
    silently bolds a blank row (and, on Sectors, aims the chart at the wrong range)."""
    ws = _book()[sheet]
    hdr = next(i for i, row in enumerate(ws.iter_rows(values_only=True), start=1)
               if row and row[0] == first_header)
    assert ws.cell(row=hdr, column=1).font.bold, f"{sheet} header row {hdr} is not bold"


def test_sector_chart_points_at_the_data_rows():
    ws = _book()["Sectors"]
    hdr = next(i for i, row in enumerate(ws.iter_rows(values_only=True), start=1)
               if row and row[0] == "sector")
    ref = str(ws._charts[0].series[0].val.numRef.f)
    # titles_from_data pulls the header into the series title, so the VALUE range covers
    # the data rows only: header+1 through header+2 for the two sectors in the fixture
    assert ref.endswith(f"$C${hdr + 1}:$C${hdr + 2}"), f"chart range {ref} is misaligned"


def test_sectors_sheet_handles_a_missing_sector():
    rows = _rows()
    rows[0]["sector"] = None
    txt = _text(_book(rows)["Sectors"])
    assert "Unclassified" in txt, "a blank sector must be labelled, not dropped silently"


# ------------------------------------------------------------------ Findings
def test_findings_sheet_leads_with_the_negative_results():
    txt = _text(_book()["Findings"])
    assert "Validated against forward returns?" in txt
    assert "FAILED three times" in txt          # E01 entry timing
    assert "SUBSTANTIALLY A SECTOR BET" in txt  # E03 R1
    assert "SCORE IS INVERTED FOR THEM" in txt  # E03 R5 cyclicals
    assert "LLM, UNREVIEWED" in txt             # the judgement half
    assert rubric.HORIZON_LABEL in txt


def test_findings_sheet_states_when_e05_has_not_been_measured(monkeypatch, tmp_path):
    monkeypatch.setattr(config, "JOURNAL_DIR", tmp_path)
    (tmp_path / "experiments").mkdir()
    txt = _text(_book()["Findings"])
    assert "NOT YET MEASURED" in txt, (
        "an unmeasured fix must not be silently omitted - omission reads as passing")
    assert "E08 - value-trap filter" in txt


def test_findings_sheet_reports_e05_verdicts_when_present(monkeypatch, tmp_path):
    import json

    monkeypatch.setattr(config, "JOURNAL_DIR", tmp_path)
    exp = tmp_path / "experiments"
    exp.mkdir()
    (exp / "E05_results.json").write_text(json.dumps({
        "f1": {"passes": False, "checks": {"within_sector_ic_positive": False,
                                           "sector_spread_halved": True}},
        "f3": {"passes": True, "checks": {"change_rate_under_12pct": True}},
        "f1_robustness": {"verdict_depends_on_sector_definition": True},
    }), encoding="utf-8")
    txt = _text(_book()["Findings"])
    assert "E05 F1 - sector-relative margins" in txt
    assert "within_sector_ic_positive=no" in txt
    assert "E05 F3 - band hysteresis" in txt


def test_findings_sheet_reports_e08_verdict_and_limit(monkeypatch, tmp_path):
    import json

    monkeypatch.setattr(config, "JOURNAL_DIR", tmp_path)
    exp = tmp_path / "experiments"
    exp.mkdir()
    (exp / "E08_results.json").write_text(json.dumps({
        "h1": {"passes": False, "checks": {"3y": False}},
        "h2": {"passes": False, "checks": {"worst": False}},
        "h3": {"passes": False, "3y": {
            "unfiltered": {"median": -0.076},
            "filtered_excluding_declining_share": {"median": -0.051},
            "share_of_top_decile_excluded": 0.04,
        }},
        "h4": {"passes": False, "checks": {"partial_ic_3y_positive": False}},
    }), encoding="utf-8")

    txt = _text(_book()["Findings"])
    assert "E08 H3" in txt and "FAIL" in txt
    assert "-7.6% unfiltered" in txt and "-5.1% after excluding 4.0%" in txt
    assert "NO_DATA names remain eligible" in txt
    assert "not true market share" in txt


# ------------------------------------------------------------------ Changes
def _write_snapshot(dirpath, day, rows):
    dirpath.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(rows).to_parquet(dirpath / f"scores_{day}.parquet", index=False)


def test_changes_sheet_says_so_when_there_is_nothing_to_compare(monkeypatch, tmp_path):
    monkeypatch.setattr(config, "JOURNAL_DIR", tmp_path)
    (tmp_path / "snapshots").mkdir()
    txt = _text(_book()["Changes"])
    assert "No earlier weekly snapshot" in txt
    assert "cannot be backfilled" in txt


def test_changes_sheet_flags_band_moves_first(monkeypatch, tmp_path):
    monkeypatch.setattr(config, "JOURNAL_DIR", tmp_path)
    old_day = (dt.date.today() - dt.timedelta(days=7)).isoformat()
    _write_snapshot(tmp_path / "snapshots", old_day, [
        {"ticker": "AAA", "composite_strict": 81, "band": "HIGH_CONVICTION",
         "coverage": 0.95},
        {"ticker": "BBB", "composite_strict": 71, "band": "INVESTABLE",
         "coverage": 0.90},   # BBB drops to WATCHLIST at 60 -> a band move
        {"ticker": "CCC", "composite_strict": 41, "band": "REJECT", "coverage": 0.85},
        {"ticker": "ZZZ", "composite_strict": 50, "band": "WEAK", "coverage": 0.70},
    ])
    ws = _book()["Changes"]
    txt = _text(ws)
    assert old_day in txt and "7 days ago" in txt

    hdr = next(i for i, row in enumerate(ws.iter_rows(values_only=True), start=1)
               if row and row[0] == "ticker")
    first = ws.cell(row=hdr + 1, column=1).value
    assert first == "BBB", "the company that changed band must sort to the top"
    assert ws.cell(row=hdr + 1, column=5).value == "YES"
    assert ws.cell(row=hdr + 1, column=8).value == -11    # 60 - 71

    assert "1 of 3 compared companies changed band" in txt
    assert "new since" in txt and "DDD" in txt      # DDD was absent from the snapshot
    assert "absent since" in txt and "ZZZ" in txt   # ZZZ has gone


def test_changes_sheet_states_the_real_gap_not_a_notional_week(monkeypatch, tmp_path):
    """A two-day comparison must not be described as week-over-week."""
    monkeypatch.setattr(config, "JOURNAL_DIR", tmp_path)
    recent = (dt.date.today() - dt.timedelta(days=2)).isoformat()
    _write_snapshot(tmp_path / "snapshots", recent,
                    [{"ticker": "AAA", "composite_strict": 80,
                      "band": "HIGH_CONVICTION", "coverage": 0.95}])
    txt = _text(_book()["Changes"])
    assert "2 days ago" in txt
    assert "weekly snapshot" not in txt


def test_changes_sheet_warns_when_the_framework_was_reweighted(monkeypatch, tmp_path):
    """A reweighting makes 500 companies move at once; without the warning that reads as
    news. This is exactly how weights v2 (FE 15->18, EN 5->2) appeared at first."""
    monkeypatch.setattr(config, "JOURNAL_DIR", tmp_path)
    day = (dt.date.today() - dt.timedelta(days=7)).isoformat()
    old = [{"ticker": r["ticker"], "composite_strict": r["composite_strict"],
            "band": r["band"], "coverage": r["coverage"],
            "fe_max": 15, "en_max": 5} for r in _rows()]
    _write_snapshot(tmp_path / "snapshots", day, old)

    rows = [dict(r, fe_max=18, en_max=2) for r in _rows()]
    txt = _text(_book(rows)["Changes"])
    assert "framework itself changed" in txt
    assert "FE 15->18" in txt and "EN 5->2" in txt


def test_changes_sheet_is_quiet_when_the_framework_did_not_change(monkeypatch, tmp_path):
    monkeypatch.setattr(config, "JOURNAL_DIR", tmp_path)
    day = (dt.date.today() - dt.timedelta(days=7)).isoformat()
    old = [{"ticker": r["ticker"], "composite_strict": r["composite_strict"],
            "band": r["band"], "coverage": r["coverage"], "fe_max": 18, "en_max": 2}
           for r in _rows()]
    _write_snapshot(tmp_path / "snapshots", day, old)
    rows = [dict(r, fe_max=18, en_max=2) for r in _rows()]
    assert "framework itself changed" not in _text(_book(rows)["Changes"])


def test_previous_snapshot_prefers_the_older_file(monkeypatch, tmp_path):
    monkeypatch.setattr(config, "JOURNAL_DIR", tmp_path)
    snaps = tmp_path / "snapshots"
    for days, score in ((10, 70), (1, 90)):
        day = (dt.date.today() - dt.timedelta(days=days)).isoformat()
        _write_snapshot(snaps, day, [{"ticker": "AAA", "composite_strict": score,
                                      "band": "X", "coverage": 0.9}])
    prev, day = xlsx_export._previous_snapshot()
    assert prev["AAA"]["composite_strict"] == 70, (
        "must skip the file written days ago and use the week-old one")
    assert day == (dt.date.today() - dt.timedelta(days=10)).isoformat()


def test_previous_snapshot_tolerates_a_corrupt_file(monkeypatch, tmp_path):
    monkeypatch.setattr(config, "JOURNAL_DIR", tmp_path)
    snaps = tmp_path / "snapshots"
    snaps.mkdir(parents=True)
    bad = (dt.date.today() - dt.timedelta(days=6)).isoformat()
    (snaps / f"scores_{bad}.parquet").write_bytes(b"not a parquet file")
    good = (dt.date.today() - dt.timedelta(days=9)).isoformat()
    _write_snapshot(snaps, good, [{"ticker": "AAA", "composite_strict": 70,
                                   "band": "X", "coverage": 0.9}])
    prev, day = xlsx_export._previous_snapshot()
    assert day == good and prev["AAA"]["composite_strict"] == 70


def test_previous_snapshot_falls_back_when_nothing_is_old_enough(monkeypatch, tmp_path):
    """Before a week of history exists the sheet should still say something true."""
    monkeypatch.setattr(config, "JOURNAL_DIR", tmp_path)
    day = (dt.date.today() - dt.timedelta(days=1)).isoformat()
    _write_snapshot(tmp_path / "snapshots", day,
                    [{"ticker": "AAA", "composite_strict": 70, "band": "X",
                      "coverage": 0.9}])
    prev, got = xlsx_export._previous_snapshot()
    assert prev is not None and got == day


# ------------------------------------------------------------------ wiring
def test_all_sheets_includes_the_new_ones_in_reading_order():
    assert xlsx_export.ALL_SHEETS[:4] == ("rank", "changes", "sectors", "findings"), (
        "the caveats must sit directly behind the ranking, not at the back")


def test_stable_export_paths_are_distinct_from_the_dated_archive():
    assert config.LATEST_XLSX.name == "company_lab_latest.xlsx"
    assert config.LATEST_CSV.name == "scores_latest.csv"
    assert config.LATEST_XLSX.parent == config.EXPORT_DIR


@pytest.mark.parametrize("sheet", ["sectors", "findings", "changes"])
def test_each_new_sheet_survives_an_empty_universe(sheet):
    wb = _book(rows=[], sheets=(sheet,))
    assert wb.sheetnames, "openpyxl refuses to save a book with no sheets"


def test_dcf_sheet_lists_valued_companies_first_then_the_refusals():
    """A DCF that produced a number for all 1,500 would be worth less than one that
    refuses 478 and says why. Both halves belong on the sheet, in that order."""
    import io

    from openpyxl import load_workbook

    from clab.export import xlsx_export

    rows = [
        {"ticker": "AAA", "name": "Cheap Co", "sector": "Industrials", "band": "WEAK",
         "selection_score": 60, "price": 10.0, "dcf_fair_value": 20.0,
         "dcf_margin_of_safety": 0.5, "dcf_growth_basis": "fcf_cagr_5y"},
        {"ticker": "BBB", "name": "Rich Co", "sector": "Industrials", "band": "WEAK",
         "selection_score": 61, "price": 40.0, "dcf_fair_value": 20.0,
         "dcf_margin_of_safety": -1.0, "dcf_growth_basis": "fcf_cagr_3y"},
        {"ticker": "CCC", "name": "A Bank", "sector": "Financials", "band": "WEAK",
         "selection_score": 55, "price": 30.0, "dcf_fair_value": None,
         "dcf_reason": "free cash flow is not a valuation input for this business"},
    ]
    data = xlsx_export.to_bytes(rows, cards=[], sheets=("dcf",))
    ws = load_workbook(io.BytesIO(data))["DCF"]
    body = [r for r in ws.iter_rows(min_row=3, values_only=True) if r[0]]
    assert [r[0] for r in body] == ["AAA", "BBB", "CCC"]      # cheapest, dearest, refused
    assert "not a valuation input" in str(body[2][-1])
