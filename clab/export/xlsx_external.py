"""The external-intelligence workbook - the human interface to everything V3 measured.

Excel is the output surface on this box; the UI is retired. Until this module existed,
every result of the external layer was visible only through a Python shell, which is the
same as not existing.

## One writer, not two

Codex writes JSON. Company Lab writes both workbooks. Two writers of one artifact cannot
be tested, and the caveats sheet would be the first thing to go missing.

## The rules borrowed from the main workbook, and why each is load-bearing

`_filterable` on every header row - **do not pre-sort and leave it inert**. The DCF
sheet shipped sorted by margin of safety with no way to ask it anything else, and that
is how a broken DCF came to rank as the cheapest company in the book.

`_finite` anywhere a column is sorted or aggregated. NaN passes
`isinstance(v, (int, float))` and does not merely skew an ordering, it scrambles it,
because every comparison against NaN is False.

Atomic tmp + `.replace()`, and `EXIT_LOCKED` when Excel holds the file open, so a
locked workbook never leaves a stale file looking current.

`Findings` sits directly behind the data on purpose. Three confident-looking conviction
scores with no measured predictive power beside them is misleading by omission.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from .. import config
from ..external import contradiction as C
from ..external import explain as EX
from ..external import gates as G
from ..external import schema as S
from ..external import score as SC
from ..external.industry_versions import active_industry_rows as select_active_industries
from ..external.store import ExternalStore
from .xlsx_export import _autosize, _filterable, _finite, _header

EXIT_LOCKED = 3

OVERVIEW_COLUMNS = [
    "conviction_rank", "sector_rank", "ticker", "name", "sector", "industry_id",
    "conviction_score", "conviction_band", "why", "why_short",
    "external_score", "external_normalized", "external_coverage",
    "external_confidence", "verified_high",
    "research_depth", "current_moat", "moat_trajectory", "competitive_position",
    "company_capture", "market_share_direction", "pricing_power",
    "regulatory_risk", "disruption_risk", "cheapness_quality",
    "flags", "label_withheld_because",
    "framework_score", "framework_band", "sg", "bq",
    "claims", "verified_claims", "distinct_domains",
    "why_cheap", "major_threat", "research_version",
]

CLAIM_COLUMNS = [
    "ticker", "industry_id", "field", "stance", "fact_or_inference",
    "independence_domain", "source_type", "verify_status", "overlap",
    "source_date", "text", "quote", "source_url", "claim_id",
]

INDUSTRY_COLUMNS = [
    "industry_id", "sector_id", "version", "structural_growth",
    "replication_difficulty", "substitution_risk", "refresh_class",
    "next_refresh_due", "key_metrics", "brief",
]


def _flat(record: dict, scored: dict, gate: dict, local: dict,
          flags: list[str], claims: list[dict]) -> dict:
    verified = sum(1 for c in claims if c.get("verify_status") == "VERIFIED_LOCAL")
    return {
        "conviction_rank": gate.get("conviction_rank"),
        "ticker": record.get("ticker"),
        "name": local.get("name"),
        "sector": local.get("sector"),
        "industry_id": record.get("industry_id"),
        "sector_rank": gate.get("sector_rank"),
        "conviction_score": gate.get("conviction_score"),
        "conviction_band": gate.get("conviction_band"),
        "why": gate.get("why"),
        "why_short": gate.get("why_short"),
        "verified_high": bool(gate.get("high_conviction_verified")),
        "external_score": scored.get("external_score"),
        "external_normalized": scored.get("external_normalized"),
        "external_coverage": scored.get("coverage"),
        "external_confidence": scored.get("external_confidence"),
        "research_depth": record.get("research_depth"),
        "current_moat": record.get("current_moat_strength"),
        "moat_trajectory": record.get("moat_trajectory"),
        "competitive_position": record.get("competitive_position"),
        "company_capture": record.get("company_specific_capture"),
        "market_share_direction": record.get("market_share_direction"),
        "pricing_power": record.get("pricing_power"),
        "regulatory_risk": record.get("regulatory_risk"),
        "disruption_risk": record.get("disruption_risk"),
        "cheapness_quality": record.get("cheapness_quality"),
        "flags": ", ".join(flags),
        "label_withheld_because": "; ".join(gate.get("disqualified_by") or []),
        "framework_score": local.get("score"),
        "framework_band": local.get("band"),
        "sg": local.get("sg"), "bq": local.get("bq"),
        "claims": len(claims), "verified_claims": verified,
        "distinct_domains": scored.get("evidence", {}).get("domains"),
        "why_cheap": record.get("why_is_it_cheap"),
        "major_threat": record.get("major_thesis_risk"),
        "research_version": record.get("research_version"),
    }



def active_industries(con) -> list[dict]:
    """The ACTIVE industry corpus: one row per industry, current version, not retired.

    `industry_intelligence` is keyed (industry_id, version) and carries a SUPERSEDED
    lifecycle, so a plain `SELECT *` returns an industry once per re-cut AND returns
    objects a re-cut retired. Both were live: `coal_consumable_fuels` had two versions,
    and E56 retired three Financials buckets whose fields are all UNKNOWN. Neither
    belongs in the weekly workbook.

    Retired operationally, retained evidentially - nothing is deleted, and the rows stay
    readable in the store and through `batch_health.object_health(include_superseded=True)`.

    Selection uses the same parsed-version helper as the store and the main workbook;
    SQL string ordering would put ``.10`` below ``.9``.
    """
    cur = con.execute("SELECT * FROM industry_intelligence")
    names = [d[0] for d in cur.description]
    rows = [dict(zip(names, r)) for r in cur.fetchall()]
    return select_active_industries(rows)


def collect(research_version: str, *, store: ExternalStore | None = None) -> dict:
    """Everything the workbook needs, computed once."""
    import pandas as pd

    store = store or ExternalStore()
    df = pd.read_parquet(config.SCORES_PARQUET).set_index("ticker")
    scored = {r["ticker"]: r for r in
              SC.score_all(research_version=research_version, store=store)}

    # Peer groups for the risk flags. Without them a risk flag cannot fire, because
    # without peers there is no way to separate a company signal from an industry
    # constant - E44 fired REGULATORY_IMPAIRMENT on 39 of 40 companies before this.
    records = store.companies(research_version)
    peers_by_industry: dict[str, list[dict]] = {}
    for rec in records:
        peers_by_industry.setdefault(rec.get("industry_id") or "", []).append(rec)

    rows, claim_rows, gates = [], [], []
    for rec in records:
        t = rec["ticker"]
        local = (df.loc[t].to_dict() if t in df.index else {})
        local["ticker"] = t
        claims = [c for c in store.claims_for(t)
                  if not rec.get("request_id")
                  or c.get("request_id") == rec["request_id"]]
        peers = [p for p in peers_by_industry.get(rec.get("industry_id") or "", [])
                 if p.get("ticker") != t]
        flags = C.flags_for(local, rec, claims, peers=peers)
        gate = G.evaluate(local, rec, scored.get(t, {}), flags)
        gate["_rec"], gate["_local"], gate["_flags"] = rec, local, flags
        gate["_claims"], gate["_scored"] = claims, scored.get(t, {})
        gate["_flags_detail"] = C.evaluate(local, rec, claims, peers=peers)
        gates.append(gate)
        claim_rows += [{k: c.get(k) for k in CLAIM_COLUMNS} for c in claims]

    ranked = G.rank(gates)
    # Rank WITHIN sector as well as overall. "These are the sectors, and these are the
    # best companies in each, top to bottom" is the question a reader actually asks, and
    # E03 measured that 95.6% of the framework's own ranking power was sector selection
    # - so a cross-sector position alone is the least informative view available.
    by_sector: dict[str, int] = {}
    for g in ranked:
        sec = (g["_local"].get("sector") or "(none)")
        by_sector[sec] = by_sector.get(sec, 0) + 1
        g["sector_rank"] = by_sector[sec]
        # Every company explains itself. A ranking that cannot say why is a black box.
        g["why"] = EX.explain(g["_rec"], g["_scored"], g, g["_flags_detail"],
                              local=g["_local"])
        g["why_short"] = EX.why_short(g, g["_flags"])
        rows.append(_flat(g["_rec"], g["_scored"], g, g["_local"],
                          g["_flags"], g["_claims"]))

    industries = []
    with store.connect() as con:
        industries = [{c: d.get(c) for c in INDUSTRY_COLUMNS}
                      for d in active_industries(con)]

    return {"overview": rows, "claims": claim_rows, "industries": industries,
            "research_version": research_version}


# ----------------------------------------------------------------- sheets
def _sheet_overview(wb, rows: list[dict]) -> None:
    ws = wb.create_sheet("Overview")
    ws.append(["Conviction is a SCORE, not a gate - every company is ranked. "
               "A withheld VERIFIED_HIGH label costs no points; see "
               "label_withheld_because."])
    _header(ws, OVERVIEW_COLUMNS)
    hdr = ws.max_row
    for r in rows:
        ws.append([r.get(c) for c in OVERVIEW_COLUMNS])
    _filterable(ws, OVERVIEW_COLUMNS, header_row=hdr)
    _autosize(ws)


def _sheet_by_sector(wb, rows: list[dict]) -> None:
    ws = wb.create_sheet("BySector")
    cols = ["sector", "sector_rank", "ticker", "conviction_score",
            "conviction_band", "current_moat", "moat_trajectory",
            "company_capture", "external_score", "external_coverage",
            "why_short", "why"]
    _header(ws, cols)
    hdr = ws.max_row
    by_sector: dict[str, list[dict]] = {}
    for r in rows:
        by_sector.setdefault(r.get("sector") or "(none)", []).append(r)
    for sector in sorted(by_sector):
        # EVERY company, top to bottom within its sector - not a top-5 cut. The point
        # is the ordering and the reason, and a truncation hides exactly the companies a
        # reader most wants to see ranked last.
        best = sorted(
            by_sector[sector],
            key=lambda x: -(x["conviction_score"]
                            if _finite(x.get("conviction_score")) else -1))
        for r in best:
            ws.append([r.get(c) for c in cols])
    _filterable(ws, cols, header_row=hdr)
    _autosize(ws)


def _sheet_claims(wb, rows: list[dict]) -> None:
    ws = wb.create_sheet("Claims")
    ws.append(["Every claim behind every field. verify_status VERIFIED_LOCAL means WE "
               "fetched the source and found the quote; SELF_ATTESTED means we could "
               "not reach it; UNVERIFIABLE means we fetched it and the quote was not "
               "there."])
    _header(ws, CLAIM_COLUMNS)
    hdr = ws.max_row
    for r in rows:
        ws.append([r.get(c) for c in CLAIM_COLUMNS])
    _filterable(ws, CLAIM_COLUMNS, header_row=hdr)
    _autosize(ws)


def _sheet_industries(wb, rows: list[dict]) -> None:
    ws = wb.create_sheet("Industries")
    _header(ws, INDUSTRY_COLUMNS)
    hdr = ws.max_row
    for r in rows:
        ws.append([r.get(c) for c in INDUSTRY_COLUMNS])
    _filterable(ws, INDUSTRY_COLUMNS, header_row=hdr)
    _autosize(ws)


def _sheet_findings(wb) -> None:
    """The caveats travel INSIDE the workbook, directly behind the data."""
    ws = wb.create_sheet("Findings")
    lines = [
        ["What this workbook does and does not establish"],
        [""],
        ["THE RANKING THIS SITS ON IS NOT VALIDATED."],
        ["E05: the top decile's median 3-year excess return over SPY was -7.6% on a "
         "survivorship-free holdout; only 45.3% beat SPY."],
        ["E03: 95.6% of the ranking's power was sector selection, IC +0.013 at 1 year."],
        ["`promoted` is false in every flag file and no code path sets it true."],
        [""],
        ["WHAT THE EXTERNAL LAYER ADDS"],
        ["It improves what is MEASURED. It is not evidence that the ranking works."],
        ["A field with no claim citing it scores NO_DATA, never zero. UNKNOWN is not "
         "the bottom of any scale."],
        ["Below 80% external coverage a company gets no external score at all."],
        [""],
        ["VALUE_TRAP_RISK COSTS ZERO POINTS, ON PURPOSE"],
        ["E08 tested the value-trap hypothesis on measured inputs and it FAILED AND "
         "INVERTED - the apparently-cheap deteriorating companies did better."],
        ["The external variant adds evidenced moat deterioration, so it is a different "
         "hypothesis, but an untested one. It is reported and never scored."],
        [""],
        ["VERIFICATION"],
        ["Claims marked VERIFIED_LOCAL were checked by fetching the source from this "
         "machine and matching the quote. A researcher re-fetching its own URLs is not "
         "verification."],
        ["Claims we could not fetch stay SELF_ATTESTED. That is our failure to check, "
         "not evidence about the company, and it is never counted as fabrication."],
        [""],
        ["E42 MEASURED THE CITATION RULE"],
        ["Arm 1 asserted 39 categoricals against 14 claims; only 10 of 36 non-UNKNOWN "
         "assertions had a claim naming that field. Its real evidenced coverage was "
         "0.31 / 0.31 / 0.15 and none of the three would have scored."],
        ["Arm 2 required a citation per field: 32 claims, 32 distinct quotes, zero "
         "reuse, 32 of 32 verified, coverage 0.85 / 0.95 / 0.77."],
    ]
    for line in lines:
        ws.append(line)
    _autosize(ws, max_width=120)


def _sheet_meta(wb, data: dict) -> None:
    ws = wb.create_sheet("Meta")
    _header(ws, ["key", "value"])
    hdr = ws.max_row
    meta = [
        ("research_version", data["research_version"]),
        ("external_schema_version", S.EXTERNAL_SCHEMA_VERSION),
        ("companies", len(data["overview"])),
        ("claims", len(data["claims"])),
        ("industry_objects", len(data["industries"])),
        ("conviction_criteria_total", G.CRITERIA_TOTAL),
        ("external_score_total", SC.TOTAL_POINTS),
        ("not_scored_fields", ", ".join(SC.UNSCORED_FIELDS)),
        ("external_min_coverage", config.EXTERNAL_MIN_COVERAGE),
        ("flag_penalties", json.dumps(G.FLAG_PENALTY)),
        ("status", "SHADOW - advisory only, promoted=false"),
    ]
    for k, v in meta:
        ws.append([k, v])
    _filterable(ws, ["key", "value"], header_row=hdr)
    _autosize(ws)


def build_workbook(data: dict):
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    _sheet_overview(wb, data["overview"])
    _sheet_by_sector(wb, data["overview"])
    _sheet_claims(wb, data["claims"])
    _sheet_industries(wb, data["industries"])
    _sheet_findings(wb)
    _sheet_meta(wb, data)
    return wb


def write(path: Path, data: dict) -> Path:
    import io

    buf = io.BytesIO()
    build_workbook(data).save(buf)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".xlsx.tmp")
    tmp.write_bytes(buf.getvalue())
    tmp.replace(path)
    return path


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Write the external-intelligence workbook")
    ap.add_argument("--version", required=True)
    ap.add_argument("--out", default=None)
    args = ap.parse_args(argv)

    data = collect(args.version)
    if not data["overview"]:
        print(f"no company records at research_version {args.version!r}")
        return 1

    out = Path(args.out) if args.out else (
        config.EXPORT_DIR / "company_lab_external_intelligence.xlsx")
    try:
        write(out, data)
    except PermissionError:
        pending = out.with_name(out.stem + "_PENDING.xlsx")
        write(pending, data)
        print(f"{out} is open in Excel; wrote {pending} instead")
        return EXIT_LOCKED

    print(f"written: {out}")
    print(f"  companies {len(data['overview'])}  claims {len(data['claims'])}  "
          f"industries {len(data['industries'])}")
    for r in data["overview"]:
        print(f"  #{r['conviction_rank']} {r['ticker']:<6} "
              f"{r['conviction_score']:>5} {r['conviction_band']:<13} "
              f"ext {r['external_score']}  {r['flags'] or '-'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
